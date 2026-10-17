"""Synthetic workbooks with known, injected problems and a ground-truth manifest.

Used by the test-suite (does each rule catch exactly the cells it should?) and by
the demo. Everything is seeded, so a given seed always produces the same file.
"""

from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

ORDER_COLUMNS = [
    "order_id", "order_date", "ship_date", "customer_id", "status", "channel",
    "quantity", "unit_price", "shipping_cost", "total", "email",
]  # fmt: skip
STATUSES = ["Pending", "Shipped", "Delivered", "Cancelled"]
CHANNELS = ["Web", "Store", "Phone"]
REGIONS = ["North", "South", "East", "West"]
N_CUSTOMERS = 40

RULES_YAML = """\
version: 1
dayfirst: true
anomalies:
  exclude: [total]          # derived column, validated by the total_matches rule
sheets:
  - sheet: Orders
    key: [order_id]
    columns:
      order_id:    {required: true, regex: "^ORD-\\\\d{5}$"}
      order_date:  {required: true, type: date, min: 2020-01-01, not_future: true}
      ship_date:   {type: date}
      customer_id: {required: true}
      status:      {required: true, allowed: [Pending, Shipped, Delivered, Cancelled]}
      quantity:    {required: true, type: int, min: 1, max: 1000}
      unit_price:  {required: true, type: float, min: 1, max: 10000}
      shipping_cost: {type: float, min: 0}
      total:       {type: float, min: 0}
      email:       {regex: "^[\\\\w.+-]+@[\\\\w-]+\\\\.[\\\\w.]+$"}
    rules:
      - id: ship_after_order
        when: "ship_date is not null"
        check: "ship_date >= order_date"
        message: "Ship date is before the order date"
      - id: total_matches
        check: "abs(total - quantity * unit_price) <= 0.01"
        message: "Total does not equal quantity x unit price"
        fix: {column: total, to: "round(quantity * unit_price, 2)"}
      - id: shipped_needs_date
        when: "status == 'Shipped'"
        require: [ship_date]
      - id: customer_exists
        lookup: {column: customer_id, sheet: Customers, key: customer_id}
  - sheet: Customers
    key: [customer_id]
    columns:
      customer_id: {required: true, regex: "^C\\\\d{3}$"}
      name:        {required: true}
      region:      {required: true, allowed: [North, South, East, West]}
      credit_limit: {type: int, min: 0}
"""


def _clean_orders(rng: random.Random, n: int) -> list[dict[str, Any]]:
    rows = []
    for i in range(1, n + 1):
        order_date = date(2024, 1, 1) + timedelta(days=rng.randint(0, 500))
        status = rng.choices(STATUSES, weights=[15, 35, 40, 10])[0]
        ship = order_date + timedelta(days=rng.randint(1, 10)) if status in ("Shipped", "Delivered") else None
        q = rng.randint(1, 50)
        price = round(rng.uniform(100, 1000), 2)
        rows.append({
            "order_id": f"ORD-{i:05d}",
            "order_date": order_date,
            "ship_date": ship,
            "customer_id": f"C{rng.randint(1, N_CUSTOMERS):03d}",
            "status": status,
            "channel": rng.choice(CHANNELS),
            "quantity": q,
            "unit_price": price,
            "shipping_cost": round(20 + 3 * q + rng.uniform(-5, 5), 2),
            "total": round(q * price, 2),
            "email": f"customer{i}@example.com",
        })
    return rows


def build_orders_workbook(path: str | Path, n: int = 300, seed: int = 7, inject: bool = True,
                          with_formula_sheet: bool = True) -> list[dict[str, Any]]:
    """Write the workbook and return the ground-truth manifest of injected problems.

    Manifest entries: {"kind": "error"|"anomaly", "sheet", "cell_ref", "rule_id"}.
    """
    rng = random.Random(seed)
    orders = _clean_orders(rng, n)
    customers = [
        {"customer_id": f"C{i:03d}", "name": f"Customer {i}", "region": rng.choice(REGIONS),
         "credit_limit": rng.choice([50_000, 100_000, 250_000])}
        for i in range(1, N_CUSTOMERS + 1)
    ]
    manifest: list[dict[str, Any]] = []
    if inject:
        manifest = _inject(orders, customers, rng)
    _write(path, orders, customers, with_formula_sheet)
    return manifest


def _inject(orders: list[dict], customers: list[dict], rng: random.Random) -> list[dict[str, Any]]:
    manifest: list[dict[str, Any]] = []
    letter = {c: get_column_letter(i + 1) for i, c in enumerate(ORDER_COLUMNS)}
    picks = iter(rng.sample(range(5, len(orders)), 40))

    def setv(**values: Any):
        return lambda r: r.update(values)

    def consistent(r: dict) -> None:
        q, p = r["quantity"], r["unit_price"]
        if isinstance(q, (int, float)) and isinstance(p, (int, float)):
            r["total"] = round(q * p, 2)
            r["shipping_cost"] = round(20 + 3 * min(q, 50) + 1.0, 2)

    def shipped(ship: date | None):
        return setv(ship_date=ship, status="Shipped" if ship else "Pending")

    def chain(*steps):
        def run(r: dict) -> None:
            for step in steps:
                step(r)
        return run

    def ship_before_order(r: dict) -> None:
        shipped(r["order_date"] - timedelta(days=3))(r)

    def total_off_by_100(r: dict) -> None:
        r["total"] = round(r["total"] + 100, 2)

    def huge_quantity(r: dict) -> None:
        r["quantity"], r["total"] = 5000, round(5000 * r["unit_price"], 2)

    # (mutation, kind, column, rule_id). Every row is used by exactly one injection.
    injections = [
        (setv(order_id=None), "error", "order_id", "required"),
        (setv(order_id="ORD-12A4"), "error", "order_id", "regex"),
        (chain(setv(order_date=date(2019, 5, 1)), shipped(date(2019, 5, 3))), "error", "order_date", "range"),
        (chain(setv(order_date=date(2031, 1, 1)), shipped(None)), "error", "order_date", "not_future"),
        (chain(setv(order_date="15/03/2024"), shipped(date(2024, 3, 18))), "error", "order_date", "type.date_as_text"),
        (chain(setv(order_date=45000), shipped(date(2023, 3, 20))), "error", "order_date", "type.date_serial"),
        (setv(order_date="not a date"), "error", "order_date", "type"),
        (ship_before_order, "error", "ship_date", "ship_after_order"),
        (setv(status="Shipped", ship_date=None), "error", "ship_date", "shipped_needs_date"),
        (setv(status="Shiped", ship_date=None), "error", "status", "allowed"),
        (setv(status="delivered"), "error", "status", "allowed.noncanonical"),
        (setv(status="Lost"), "error", "status", "allowed"),
        (setv(status="Pending ", ship_date=None), "error", "status", "format.whitespace"),
        (setv(quantity="abc"), "error", "quantity", "type"),
        (chain(setv(quantity=12), consistent, setv(quantity="12")), "error", "quantity", "type.stored_as_text"),
        (chain(setv(quantity=0), consistent), "error", "quantity", "range"),
        (huge_quantity, "error", "quantity", "range"),
        (chain(setv(unit_price=1200.5), consistent, setv(unit_price="1,200.50")), "error", "unit_price",
         "type.stored_as_text"),
        (total_off_by_100, "error", "total", "total_matches"),
        (setv(customer_id=None), "error", "customer_id", "required"),
        (setv(customer_id="C999"), "error", "customer_id", "customer_exists"),
        (setv(email="not-an-email"), "error", "email", "regex"),
        # Anomalies: suspicious but rule-compliant values.
        (chain(setv(quantity=1, unit_price=9800.0), consistent), "anomaly", "unit_price", "anomaly.robust_z"),
        (chain(setv(quantity=3), consistent, setv(shipping_cost=950.0)), "anomaly", "shipping_cost",
         "anomaly.robust_z"),
        # Within each column's normal range, but shipping cost no longer fits the quantity.
        (chain(setv(quantity=2), consistent, setv(shipping_cost=165.0)), "anomaly", "shipping_cost",
         "anomaly.isolation_forest"),
        (chain(setv(quantity=49), consistent, setv(shipping_cost=22.0)), "anomaly", "quantity",
         "anomaly.isolation_forest"),
        (setv(channel="Fax"), "anomaly", "channel", "anomaly.rare_category"),
    ]
    for mutate, kind, col, rule_id in injections:
        i = next(picks)
        mutate(orders[i])
        manifest.append({"kind": kind, "sheet": "Orders", "cell_ref": f"Orders!{letter[col]}{i + 2}",
                         "rule_id": rule_id})

    # Exact duplicate of an untouched row, appended at the end.
    src = next(picks)
    orders.append(dict(orders[src]))
    row = len(orders) + 1  # Excel row of the appended copy
    manifest.append({"kind": "error", "sheet": "Orders", "cell_ref": f"Orders!A{row}", "rule_id": "row.duplicate"})
    manifest.append({"kind": "error", "sheet": "Orders", "cell_ref": f"Orders!{letter['order_id']}{row}",
                     "rule_id": "key.duplicate"})

    # Customers sheet: one bad region, one region with odd casing.
    customers[4]["region"] = "Nort"
    manifest.append({"kind": "error", "sheet": "Customers", "cell_ref": "Customers!C6", "rule_id": "allowed"})
    customers[9]["region"] = "south"
    manifest.append({"kind": "error", "sheet": "Customers", "cell_ref": "Customers!C11", "rule_id": "allowed.noncanonical"})
    return manifest


def _write(path: str | Path, orders: list[dict], customers: list[dict], with_formula_sheet: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(ORDER_COLUMNS)
    for r in orders:
        ws.append([r[c] for c in ORDER_COLUMNS])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):
        for c in row:
            if isinstance(c.value, date):
                c.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    wc = wb.create_sheet("Customers")
    cols = ["customer_id", "name", "region", "credit_limit"]
    wc.append(cols)
    for r in customers:
        wc.append([r[c] for c in cols])
    if with_formula_sheet:
        s = wb.create_sheet("Summary")
        s.append(["metric", "value"])
        s.append(["total_revenue", "=SUM(Orders!J:J)"])
        s.append(["orders", "=COUNTA(Orders!A:A)-1"])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_edge_workbook(path: str | Path) -> None:
    """Title row, merged header, blank row, leading-zero ids, serial date, formulas, empty sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Quarterly Export"
    ws.merge_cells("A1:F1")
    ws.append([])  # row 2 blank
    ws.append(["ID", "Name", "Amount", None, "Joined", "Double"])  # row 3
    ws.merge_cells("C3:D3")
    ws.append(["00123", "Alpha", 10, 1, date(2024, 1, 5), "=C4*2"])  # row 4
    ws.append(["00124", "Beta ", 20, 2, 45000, "=C5*2"])  # row 5
    ws.append([None, None, None, None, None, None])  # row 6 blank
    ws.append(["00125", "Gamma", 30, 3, "2024-02-01", "=C7*2"])  # row 7
    wb.create_sheet("Empty")
    wb.save(path)
