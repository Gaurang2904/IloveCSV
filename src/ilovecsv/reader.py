"""Load .xlsx / .csv files into DataFrames that remember where every cell came from.

Each sheet becomes an object-dtype DataFrame whose index is the Excel row
number and whose columns map back to Excel column letters, so every issue can
point at a real cell such as ``Orders!D14``. Values are kept raw (no type
coercion), so IDs like ``00123`` survive and the profiler sees what is really in
the file.
"""

from __future__ import annotations

import csv
import hashlib
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

NATIVE = {".xlsx", ".xlsm", ".csv", ".tsv", ".txt"}  # read here; other formats go through adapters
HEADER_SCAN_ROWS = 30


class ReadError(Exception):
    """The file cannot be read as a dataset."""


class LargeDatasetError(ReadError):
    """The table is larger than the full-processing limit; inspect it, or process a sample."""


@dataclass
class SheetData:
    name: str
    df: pd.DataFrame  # index = Excel row number, dtype=object
    header_row: int
    col_letters: dict[str, str]
    formula_cells: set[str] = field(default_factory=set)  # coordinates such as "H5"
    dropped_blank_rows: list[int] = field(default_factory=list)
    dropped_blank_cols: list[str] = field(default_factory=list)

    def coord(self, xl_row: int, column: str) -> str:
        return f"{self.col_letters[column]}{xl_row}"

    def cell_ref(self, xl_row: int, column: str) -> str:
        return f"{self.name}!{self.coord(xl_row, column)}"

    def row_ref(self, xl_row: int) -> str:
        first = next(iter(self.col_letters.values()), "A")
        return f"{self.name}!{first}{xl_row}"

    def is_formula(self, xl_row: int, column: str) -> bool:
        return self.coord(xl_row, column) in self.formula_cells


@dataclass
class WorkbookData:
    """A source file as one or more named tables ("sheets"), whatever its original format."""

    filename: str
    path: Path
    file_type: str  # "xlsx" | "csv" (cell-level write-back) | "frames" (any other adapter)
    sha256: str
    sheets: dict[str, SheetData]
    warnings: list[str] = field(default_factory=list)
    csv_delimiter: str | None = None
    source_format: str = ""  # e.g. "xlsx", "csv", "jsonl", "parquet", "zip", "mat", "nasa_battery_mat"
    tables_meta: dict[str, dict[str, Any]] = field(default_factory=dict)  # per table: source, conversion notes
    docs: list[dict[str, Any]] = field(default_factory=list)  # documentation found next to the data
    unsupported: list[dict[str, Any]] = field(default_factory=list)  # readable parts with no safe table form
    partial: dict[str, Any] | None = None  # set when only a sample of rows was processed
    encoding: str | None = None  # text sources
    tree: list[dict[str, Any]] = field(default_factory=list)  # hierarchical sources: what is inside

    @property
    def first_sheet(self) -> SheetData:
        return next(iter(self.sheets.values()))


def native(v: Any) -> Any:
    """Plain Python value for the engine: numpy scalars, Timestamps, NaN/NaT, bytes, nested data."""
    if v is None or v is pd.NaT:
        return None
    if isinstance(v, float):
        return None if v != v else v
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().replace(tzinfo=None) if v.tzinfo else v.to_pydatetime()
    if isinstance(v, (bytes, bytearray)):
        return v.decode("utf-8", errors="replace")
    if hasattr(v, "dtype") and hasattr(v, "item") and getattr(v, "ndim", 0) == 0:  # numpy scalar
        if str(v.dtype).startswith("datetime64"):
            ts = pd.Timestamp(v)
            return None if ts is pd.NaT else ts.to_pydatetime()
        return native(v.item())
    if isinstance(v, (list, tuple, dict)) or (hasattr(v, "tolist") and getattr(v, "ndim", 0) > 0):
        import json

        data = v.tolist() if hasattr(v, "tolist") else v
        return json.dumps(data, ensure_ascii=False, default=str)
    return v


def sheet_from_frame(name: str, df: pd.DataFrame) -> SheetData:
    """Build a table from a DataFrame, numbered like a sheet (header row 1, data from row 2)."""
    names: list[str] = []
    seen: dict[str, int] = {}
    for c in df.columns:
        base = str(c).strip() or "column"
        if base in seen:
            seen[base] += 1
            base = f"{base}.{seen[base]}"
        else:
            seen[base] = 0
        names.append(base)
    data: dict[str, list[Any]] = {}
    for new, col in zip(names, df.columns, strict=True):
        s = df[col]
        if pd.api.types.is_datetime64_any_dtype(s):
            data[new] = [None if pd.isna(x) else native(pd.Timestamp(x)) for x in s]
        else:
            data[new] = [native(x) for x in s.astype(object).tolist()]
    out = pd.DataFrame(data, columns=names, dtype=object)
    out.index = pd.Index(range(2, len(out) + 2), name="xl_row")
    letters = {n: get_column_letter(i + 1) for i, n in enumerate(names)}
    return SheetData(name=name[:100], df=out, header_row=1, col_letters=letters)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_workbook(path: str | Path, filename: str | None = None, options: Any = None) -> WorkbookData:
    """Read any supported file into tables. `options` is a sources.SourceOptions (or None)."""
    from .sources import SourceOptions

    options = options or SourceOptions()
    path = Path(path)
    filename = filename or path.name
    suffix = Path(filename).suffix.lower()
    if suffix not in NATIVE:
        from .adapters import load_dataset  # JSON, Parquet, ZIP, MATLAB, NumPy, HDF5, ...

        return load_dataset(path, filename, options)
    digest = sha256_file(path)
    if suffix in (".csv", ".tsv", ".txt"):
        sheet, delimiter, partial, encoding = _read_csv(path, Path(filename).stem, options)
        wb = WorkbookData(filename, path, "csv", digest, {sheet.name: sheet}, csv_delimiter=delimiter,
                          source_format=suffix.lstrip("."), partial=partial, encoding=encoding)
        if partial:
            wb.warnings.append(f"Only the first {partial['rows']} rows were processed (max_rows).")
        return wb
    wb = _read_xlsx(path, filename, digest)
    wb.source_format = suffix.lstrip(".")
    if suffix == ".xlsm":
        wb.warnings.append("Macro-enabled workbook: macros were not run, and the cleaned copy is saved as .xlsx "
                           "without them.")
    if options.sheets:
        missing = [s for s in options.sheets if s not in wb.sheets]
        if missing:
            raise ReadError(f"Sheet(s) not found: {missing}. Available: {list(wb.sheets)}")
        wb.sheets = {k: v for k, v in wb.sheets.items() if k in options.sheets}
    for name, sheet in wb.sheets.items():
        enforce_row_limit(name, len(sheet.df), options)
    return wb


def read_table_csvs(files: dict[str, Path], original: WorkbookData) -> WorkbookData:
    """Re-read cleaned per-table CSVs (non-spreadsheet sources) under their table names, so the
    exported files themselves are what gets re-validated."""
    from .sources import SourceOptions

    opts = SourceOptions(delimiter=",", encoding="utf-8", row_limit=10**12)
    sheets: dict[str, SheetData] = {}
    for table, path in files.items():
        sheet, _, _, _ = _read_csv(Path(path), table, opts)
        sheet.name = table
        sheets[table] = sheet
    return WorkbookData(original.filename, Path(next(iter(files.values()))), "csv",
                        sha256_file(next(iter(files.values()))), sheets, csv_delimiter=",",
                        source_format="csv", encoding="utf-8")


def text_values(wb: WorkbookData, table: str) -> str | None:
    """Which values a table's source can only hold as text, so text there is not a problem:
    "all" (CSV/TSV, also inside archives), "dates" (JSON has no date type), or None (typed formats)."""
    if wb.file_type == "csv":
        return "all"
    return (wb.tables_meta.get(table) or {}).get("text_values")


def enforce_row_limit(table: str, rows: int, options: Any) -> None:
    if options.max_rows is None and rows > options.row_limit:
        raise LargeDatasetError(
            f"Table '{table}' has {rows:,} rows, above the full-processing limit of {options.row_limit:,}. "
            "Inspect it first (ilovecsv inspect), or process a sample with max_rows.")


# --------------------------------------------------------------------------- xlsx


def _read_xlsx(path: Path, filename: str, digest: str) -> WorkbookData:
    warnings: list[str] = []
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
    except zipfile.BadZipFile as e:
        raise ReadError("The file is not a valid .xlsx workbook.") from e
    if any(n.startswith(("xl/charts/", "xl/drawings/")) for n in names):
        warnings.append(
            "The workbook contains charts or drawings; openpyxl may not preserve them "
            "in the corrected copy."
        )
    try:
        wb_values = load_workbook(path, data_only=True)
        wb_formulas = load_workbook(path, data_only=False)
    except Exception as e:  # openpyxl raises many types for corrupt files
        raise ReadError(f"Could not open workbook: {e}") from e

    sheets: dict[str, SheetData] = {}
    for ws in wb_values.worksheets:
        ws_f = wb_formulas[ws.title]
        formula_cells = {
            c.coordinate for row in ws_f.iter_rows() for c in row if c.data_type == "f"
        }
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        merged = [
            (m.min_row, m.max_row, m.min_col, m.max_col) for m in ws.merged_cells.ranges
        ]
        sheet = _grid_to_sheet(ws.title, grid, merged, formula_cells)
        if sheet is None:
            warnings.append(f"Sheet '{ws.title}' is empty and was skipped.")
            continue
        sheets[sheet.name] = sheet
    if not sheets:
        raise ReadError("The workbook has no sheets with data.")
    return WorkbookData(filename, path, "xlsx", digest, sheets, warnings)


# --------------------------------------------------------------------------- csv


def detect_text_encoding(head: bytes, preferred: str | None = None) -> str:
    candidates = [preferred] if preferred else ["utf-8-sig", "cp1252"]
    for enc in candidates:
        try:
            head.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
        except LookupError as e:
            raise ReadError(f"Unknown encoding '{enc}'.") from e
    if preferred:
        raise ReadError(f"The file is not valid {preferred} text.")
    return "cp1252"  # pragma: no cover - cp1252 decodes almost anything


def count_lines(path: Path) -> int:
    """Physical line count without decoding (fast, for size estimates)."""
    n, last = 0, b""
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 22), b""):
            n += chunk.count(b"\n")
            last = chunk[-1:]
    return n + (1 if last and last != b"\n" else 0)


def _read_csv(path: Path, stem: str, options: Any = None) -> tuple[SheetData, str, dict[str, Any] | None, str]:
    from .sources import SourceOptions

    options = options or SourceOptions()
    with open(path, "rb") as f:
        head = f.read(1 << 16)
    if b"\x00" in head[:4096]:
        raise ReadError("The CSV file contains binary data.")
    encoding = detect_text_encoding(head, options.encoding)
    sample = head.decode(encoding, errors="replace")
    delimiter = options.delimiter
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    lines = count_lines(path)
    if options.max_rows is None:
        enforce_row_limit(stem, max(0, lines - 1), options)
    limit = None if options.max_rows is None else options.max_rows + HEADER_SCAN_ROWS
    grid: list[list[Any]] = []
    with open(path, encoding=encoding, newline="", errors="strict" if options.encoding else "replace") as f:
        for row in csv.reader(f, delimiter=delimiter):
            grid.append([None if c == "" else c for c in row])
            if limit is not None and len(grid) >= limit:
                break
    name = (stem or "data")[:31]
    sheet = _grid_to_sheet(name, grid, [], set())
    if sheet is None:
        raise ReadError("The CSV file has no data.")
    partial = None
    if options.max_rows is not None and (len(sheet.df) > options.max_rows or limit and len(grid) >= limit):
        sheet.df = sheet.df.iloc[: options.max_rows]
        partial = {"rows": len(sheet.df), "estimated_total_rows": max(0, lines - sheet.header_row)}
    return sheet, delimiter, partial, encoding


# --------------------------------------------------------------------------- shared


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _detect_header(grid: list[list[Any]]) -> int | None:
    scan = grid[:HEADER_SCAN_ROWS]
    widths = [sum(not _blank(v) for v in r) for r in scan]
    if not widths or max(widths) == 0:
        return None
    max_width = max(widths)
    for i, row in enumerate(scan):
        filled = [v for v in row if not _blank(v)]
        if len(filled) < max(1, 0.5 * max_width):
            continue
        text = sum(isinstance(v, str) for v in filled)
        if text / len(filled) >= 0.5:
            return i
    return next(i for i, w in enumerate(widths) if w)


def _grid_to_sheet(
    name: str,
    grid: list[list[Any]],
    merged: list[tuple[int, int, int, int]],
    formula_cells: set[str],
) -> SheetData | None:
    width = max((len(r) for r in grid), default=0)
    grid = [r + [None] * (width - len(r)) for r in grid]
    h = _detect_header(grid)
    if h is None:
        return None
    header_xl_row = h + 1

    # Merged header cells keep their value only in the top-left cell: spread it.
    header = list(grid[h])
    for min_r, max_r, min_c, max_c in merged:
        if min_r <= header_xl_row <= max_r:
            top_left = grid[min_r - 1][min_c - 1]
            for c in range(min_c - 1, max_c):
                if _blank(header[c]):
                    header[c] = top_left

    body = grid[h + 1 :]
    # Drop trailing blank rows (formatting often extends the used range).
    while body and all(_blank(v) for v in body[-1]):
        body.pop()

    keep_cols: list[int] = []
    dropped_cols: list[str] = []
    for c in range(width):
        if _blank(header[c]) and all(_blank(r[c]) for r in body):
            if c < max((i for i, v in enumerate(header) if not _blank(v)), default=-1):
                dropped_cols.append(get_column_letter(c + 1))
            continue
        keep_cols.append(c)

    names: list[str] = []
    seen: dict[str, int] = {}
    for c in keep_cols:
        v = header[c]
        base = str(v).strip() if not _blank(v) else f"column_{get_column_letter(c + 1)}"
        if base in seen:
            seen[base] += 1
            base = f"{base}.{seen[base]}"
        else:
            seen[base] = 0
        names.append(base)

    rows: list[list[Any]] = []
    index: list[int] = []
    dropped_rows: list[int] = []
    for offset, r in enumerate(body):
        xl_row = header_xl_row + 1 + offset
        values = [r[c] for c in keep_cols]
        if all(_blank(v) for v in values):
            dropped_rows.append(xl_row)
            continue
        rows.append(values)
        index.append(xl_row)

    df = pd.DataFrame(rows, columns=names, index=pd.Index(index, name="xl_row"), dtype=object)
    if df.empty and not names:
        return None
    letters = {n: get_column_letter(c + 1) for n, c in zip(names, keep_cols, strict=True)}
    return SheetData(
        name=name,
        df=df,
        header_row=header_xl_row,
        col_letters=letters,
        formula_cells=formula_cells,
        dropped_blank_rows=dropped_rows,
        dropped_blank_cols=dropped_cols,
    )
