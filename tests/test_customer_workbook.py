"""A noisy customer workbook (the shape of a typical CRM export) from upload to corrected file.

Covers the problems a reviewer expects a cleaning tool to handle: duplicate records that differ
only in their generated id, missing values, malformed emails and phones, impossible ages,
inconsistent capitalisation, a known city alias, and a Readme sheet that is not data.
"""

from datetime import date

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from ilovecsv.pipeline import decide, export, run_checks
from ilovecsv.profiling import profile_workbook
from ilovecsv.reader import read_workbook
from ilovecsv.rules.suggest import suggest_rules

HEADER = ["Customer ID", "Full Name", "Email", "Phone", "City", "Age", "Signup Date", "Plan", "Monthly Spend", "Status"]
CITIES = ["Pune", "Mumbai", "Delhi", "Chennai", "Kolkata", "Bengaluru", "Jaipur", "Surat"]
PLANS = [("Basic", 799), ("Standard", 1299), ("Premium", 2499)]


def build(path):
    rows = []
    for i in range(1, 25):
        plan, spend = PLANS[i % 3]
        rows.append([f"K{i:03d}", f"Person {chr(64 + i)} Test", f"person{i}@example.org", f"98200{i:05d}",
                     CITIES[i % len(CITIES)], 25 + (i * 7) % 20, date(2025, 1 + i % 12, 1 + i % 27), plan, spend,
                     "Active" if i % 4 else "Inactive"])
    rows[1][4] = "pune"                    # K002 case variant        -> Pune
    rows[9][4] = "PUNE"                    # K010 case variant        -> Pune
    rows[2][4] = "Bangalore"               # K003 alias (rule)        -> Bengaluru
    rows[3][3] = "+91 98200 00004"         # K004 phone style         -> 9820000004
    rows[4][3] = "98200-00005"             # K005 phone style         -> 9820000005
    rows[5][2] = "PERSON6@Example.org"     # K006 email case          -> person6@example.org
    rows[6][2] = "person7.example.org"     # K007 invalid email       (human)
    rows[7][2] = None                      # K008 missing email       (human)
    rows[8][5] = None                      # K009 missing age         (human)
    rows[10][5] = -2                       # K011 impossible age      (human)
    rows[11][5] = 140                      # K012 impossible age      (human)
    rows[12][6] = "2025-02-30"             # K013 impossible date     (human)
    rows[13][7] = "premium"                # K014 plan casing         -> Premium
    dup = list(rows[14])                   # K015 duplicated as K025, only the id differs
    dup[0] = "K025"
    rows.append(dup)
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    notes = wb.create_sheet("Readme")
    notes.append(["Dataset Notes"])
    notes.append(["Synthetic test data. No real people are represented in this workbook at all."])
    notes.append(["It intentionally contains duplicates, missing values and inconsistent formatting."])
    wb.save(path)
    return path


@pytest.fixture(scope="module")
def flow(tmp_path_factory):
    d = tmp_path_factory.mktemp("cust")
    path = build(d / "customers.xlsx")
    doc = yaml.safe_load(suggest_rules(profile_workbook(read_workbook(path))))
    doc["columns"]["City"]["aliases"] = {"Bangalore": "Bengaluru"}  # reference knowledge a person adds
    rules = yaml.safe_dump(doc, sort_keys=False)
    run = run_checks(path, rules)
    before = run.summary()
    decide(run, min_confidence=0.9, approver="tester")
    ex = export(run, d / "out")
    return path, doc, run, before, ex


def found(run):
    return {(i.rule_id, i.cell_ref) for i in run.validation.issues}


def test_suggested_rules(flow):
    _, doc, *_ = flow
    assert "sheets" not in doc and doc["sheet"] == "Customers"  # Readme skipped
    cols = doc["columns"]
    assert cols["Email"]["format"] == "email" and cols["Email"]["required"]
    assert cols["Phone"]["format"] == "phone" and cols["Phone"]["phone_digits"] == 10
    assert cols["Age"]["required"] and cols["Age"]["min"] == 0 and cols["Age"]["max"] < 140
    assert cols["Signup Date"]["type"] == "date"


def test_detects_every_planted_problem(flow):
    _, _, run, before, _ = flow
    # City has few distinct values here, so it is a category with an allowed list and the
    # case variants surface as allowed.noncanonical (format.case_variant for free-text columns).
    issues = found(run)
    for cell in ("Customers!E3", "Customers!E11"):
        assert {("allowed.noncanonical", cell), ("format.case_variant", cell)} & issues, cell
    expected = {
        ("row.duplicate", "Customers!A26"),
        ("allowed.alias", "Customers!E4"),
        ("format.phone_style", "Customers!D5"), ("format.phone_style", "Customers!D6"),
        ("format.email_case", "Customers!C7"), ("format.email", "Customers!C8"),
        ("required", "Customers!C9"), ("required", "Customers!F10"),
        ("range", "Customers!F12"), ("range", "Customers!F13"),
        ("type", "Customers!G14"),
        ("allowed.noncanonical", "Customers!H15"),
    }
    assert expected <= found(run), sorted(expected - found(run))
    assert before["rows_by_sheet"] == {"Customers": 25}


def test_export_applies_approved_fixes(flow):
    path, _, run, before, ex = flow
    assert ex.apply.applied and not ex.apply.skipped
    out = load_workbook(ex.corrected_path)["Customers"]
    values = {row[0].value: [c.value for c in row] for row in out.iter_rows(min_row=2)}
    assert "K025" not in values and len(values) == 24                 # duplicate removed
    assert values["K002"][4] == "Pune" and values["K010"][4] == "Pune"
    assert values["K003"][4] == "Bengaluru"
    assert values["K004"][3] == "9820000004" and values["K005"][3] == "9820000005"
    assert values["K006"][2] == "person6@example.org"
    assert values["K014"][7] == "Premium"
    assert values["K007"][2] == "person7.example.org"                # never guessed
    assert values["K009"][5] is None and values["K011"][5] == -2     # never imputed
    assert load_workbook(ex.corrected_path)["Readme"].max_row == 3


def test_score_and_remaining_errors(flow):
    _, _, run, before, ex = flow
    after = ex.after
    assert after["quality_score"] > before["quality_score"]
    assert set(after["by_rule"]) == {"format.email", "required", "range", "type"}  # only what needs a human
    assert after["errors"] == 6 and after["warnings"] == 0


def test_export_without_approvals_is_flagged(tmp_path):
    path = build(tmp_path / "c.xlsx")
    rules = suggest_rules(profile_workbook(read_workbook(path)))
    run = run_checks(path, rules)
    ex = export(run, tmp_path / "out")  # nothing approved
    assert ex.apply.applied == []
    html = ex.reports["html"].read_text(encoding="utf-8")
    assert "No changes were applied" in html
