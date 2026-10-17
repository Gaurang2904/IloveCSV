import json

from ilovecsv.pipeline import decide, export, run_checks
from ilovecsv.reader import sha256_file
from ilovecsv.rules.schema import load_rules, rules_sha256
from ilovecsv.rules.suggest import suggest_rules


def test_same_input_same_issues(orders, rules_text):
    a = run_checks(orders[0], rules_text)
    b = run_checks(orders[0], rules_text)
    assert [i.to_dict() for i in a.issues] == [i.to_dict() for i in b.issues]
    assert [s.to_dict() for s in a.suggestions] == [s.to_dict() for s in b.suggestions]


def test_reports_carry_provenance(orders, rules_text, tmp_path):
    run = run_checks(orders[0], rules_text)
    decide(run, min_confidence=0.9, approver="auditor")
    ex = export(run, tmp_path)
    audit = json.loads(ex.reports["json"].read_text(encoding="utf-8"))
    assert audit["input"]["sha256"] == sha256_file(orders[0])
    assert audit["rules"]["sha256"] == rules_sha256(rules_text)
    assert audit["output"]["sha256"] == sha256_file(ex.corrected_path)
    assert audit["changes"] and all(c["approved_by"] == "auditor" for c in audit["changes"])
    assert audit["after"]["errors"] is not None
    assert len(audit["issues"]) == len(run.validation.issues)
    assert all(i["why"] for i in audit["issues"])
    html = ex.reports["html"].read_text(encoding="utf-8")
    assert "Data Quality Report" in html and audit["input"]["sha256"] in html
    from openpyxl import load_workbook

    sheets = load_workbook(ex.reports["xlsx"]).sheetnames
    assert sheets == ["Summary", "Profile", "Issues", "Anomalies", "Changes", "Skipped"]


def test_suggested_rules_are_valid_and_pass_on_clean_data(clean_orders):
    from ilovecsv.profiling import profile_workbook
    from ilovecsv.reader import read_workbook

    text = suggest_rules(profile_workbook(read_workbook(clean_orders)))
    rules = load_rules(text)
    assert {s.sheet for s in rules.sheets} >= {"Orders", "Customers"}
    orders = next(s for s in rules.sheets if s.sheet == "Orders")
    assert orders.key == ["order_id"]
    assert set(orders.columns["status"].allowed) == {"Pending", "Shipped", "Delivered", "Cancelled"}
    run = run_checks(clean_orders, text)
    assert run.summary()["errors"] == 0
