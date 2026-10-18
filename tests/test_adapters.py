"""Format adapters: every format becomes clearly defined tables, and what cannot is reported."""

from __future__ import annotations

import json
import zipfile

import numpy as np
import pandas as pd
import pytest
import yaml

from ilovecsv.adapters import inspect_file, signature_problem
from ilovecsv.cli import main as cli
from ilovecsv.pipeline import prepare
from ilovecsv.reader import LargeDatasetError, ReadError, read_workbook
from ilovecsv.sources import SourceOptions

FRAME = pd.DataFrame({"id": [1, 2, 3, 4], "temp": [20.5, 21.0, -999.0, 22.1], "label": ["a", "b", "a", "b"]})


def tables(path, **opts):
    wb = read_workbook(path, path.name, SourceOptions(**opts))
    return wb, {k: s.df for k, s in wb.sheets.items()}


# ----------------------------------------------------------------------------- tier 1
@pytest.mark.parametrize("ext", ["parquet", "feather"])
def test_columnar_formats(tmp_path, ext):
    path = tmp_path / f"t.{ext}"
    getattr(FRAME, f"to_{ext}")(path)
    wb, t = tables(path)
    assert wb.file_type == "frames" and wb.source_format == ext
    assert list(t["t"].columns) == ["id", "temp", "label"] and len(t["t"]) == 4
    insp = inspect_file(path)
    assert insp.tables[0].rows == 4 and insp.tables[0].rows_exact


def test_jsonl_and_nested_json(tmp_path):
    jl = tmp_path / "events.jsonl"
    jl.write_text("\n".join(json.dumps(r) for r in FRAME.to_dict("records")))
    _, t = tables(jl)
    assert len(t["events"]) == 4
    nested = tmp_path / "orders.json"
    nested.write_text(json.dumps({"meta": {"n": 2}, "data": {"items": [
        {"id": 1, "customer": {"name": "Ann", "city": "Pune"}, "tags": ["x", "y"]},
        {"id": 2, "customer": {"name": "Raj", "city": "Mumbai"}, "tags": []}]}}))
    wb, t = tables(nested, json_records="data.items")
    df = next(iter(t.values()))
    assert {"customer.name", "customer.city"} <= set(df.columns) and len(df) == 2
    auto_wb, auto = tables(nested)  # the only array of records is found without a path
    assert len(next(iter(auto.values()))) == 2
    with pytest.raises(ReadError):
        tables(nested, json_records="data.missing")


def test_arff_missing_marker(tmp_path):
    path = tmp_path / "w.arff"
    path.write_text("@relation w\n@attribute id numeric\n@attribute temp numeric\n@attribute label {a,b}\n"
                    "@data\n1,20.5,a\n2,?,b\n3,22.1,a\n")
    _, t = tables(path)
    df = t["w"]
    assert len(df) == 3
    assert df["temp"].iloc[1] is None or pd.isna(df["temp"].iloc[1])


def test_xlsm_read_with_warning(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "book.xlsm"
    wb = Workbook()
    wb.active.append(["id", "score"])
    wb.active.append([1, 2.5])
    wb.save(path)
    data = read_workbook(path)
    assert data.file_type == "xlsx"
    assert any("macro" in w.lower() for w in data.warnings)


def test_xls_legacy(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    path = tmp_path / "old.xls"
    book = xlwt.Workbook()
    sh = book.add_sheet("data")
    for j, v in enumerate(["id", "score"]):
        sh.write(0, j, v)
    sh.write(1, 0, 1)
    sh.write(1, 1, 2.5)
    book.save(str(path))
    _, t = tables(path)
    assert list(t["data"].columns) == ["id", "score"]


def test_signature_mismatch(tmp_path):
    fake = tmp_path / "x.parquet"
    fake.write_text("id,v\n1,2\n")
    assert signature_problem(fake, fake.name)
    binary = tmp_path / "x.csv"
    binary.write_bytes(b"PK\x03\x04 not text")
    assert signature_problem(binary, binary.name)
    ok = tmp_path / "ok.csv"
    ok.write_text("a,b\n1,2\n")
    assert signature_problem(ok, ok.name) is None


# ----------------------------------------------------------------------------- ZIP
def test_zip_tables_docs_and_merge(tmp_path):
    path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("2023/sensors.csv", "id,v\n1,2\n2,-999\n")
        z.writestr("2024/sensors.csv", "id,v\n3,4\n")
        z.writestr("labels.json", json.dumps([{"id": 1, "y": "ok"}]))
        z.writestr("README.md", "Column v: voltage. Missing values are coded as -999.\n")
        z.writestr("notes.bin", b"\x00\x01")
    wb, t = tables(path)
    assert set(t) == {"sensors", "labels"}
    assert len(t["sensors"]) == 3 and "source_file" in t["sensors"].columns  # same name + same columns: merged
    assert any(h["token"] == "-999" for d in wb.docs for h in d["missing_value_hints"])
    insp = inspect_file(path)
    assert {x.name for x in insp.tables} == {"sensors", "labels"}
    _, separate = tables(path, combine_tables=False)
    assert set(separate) == {"2023_sensors", "2024_sensors", "labels"}


def test_zip_refuses_unsafe_paths(tmp_path):
    path = tmp_path / "evil.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("../../escape.csv", "a\n1\n")
    with pytest.raises(ReadError, match="unsafe path"):
        read_workbook(path)
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("ok.csv", "a\n1\n")
    wb = read_workbook(path)
    assert set(wb.sheets) == {"ok"}
    assert any("unsafe" in u["reason"] for u in wb.unsupported)
    assert not (tmp_path.parent / "escape.csv").exists()


# ----------------------------------------------------------------------------- tier 2
def test_numpy(tmp_path):
    np.save(tmp_path / "a.npy", np.arange(12).reshape(4, 3))
    _, t = tables(tmp_path / "a.npy")
    assert list(t["a"].columns) == ["c0", "c1", "c2"] and len(t["a"]) == 4
    np.save(tmp_path / "cube.npy", np.zeros((2, 2, 2)))
    insp = inspect_file(tmp_path / "cube.npy")
    assert not insp.tables and "3-D" in insp.unsupported[0]["reason"]
    np.savez(tmp_path / "b.npz", x=np.arange(6).reshape(3, 2), y=np.array([0, 1, 0]))
    _, t = tables(tmp_path / "b.npz")
    assert set(t) == {"x", "y"}
    np.save(tmp_path / "obj.npy", np.array([{"a": 1}], dtype=object), allow_pickle=True)
    with pytest.raises(ReadError):  # pickled objects are never loaded
        read_workbook(tmp_path / "obj.npy")


def test_mat_generic_reports_what_is_not_a_table(tmp_path):
    from scipy.io import savemat

    path = tmp_path / "m.mat"
    savemat(path, {"signal": np.arange(10.0).reshape(5, 2), "fs": 100.0, "names": np.array(["ab", "cd"]),
                   "cube": np.zeros((2, 2, 2))})
    wb, t = tables(path)
    assert t["signal"].shape == (5, 2) and len(t["names"]) == 2
    reasons = {u["path"]: u["reason"] for u in wb.unsupported}
    assert "fs" in reasons and "cube" in reasons and "3-D" in reasons["cube"]


def _nasa_like(path):
    from scipy.io import savemat

    dt = [("type", "O"), ("ambient_temperature", "O"), ("time", "O"), ("data", "O")]
    cyc = np.empty((1, 3), dtype=dt)
    clock = np.array([2008.0, 4, 2, 13, 8, 17.9])
    cyc[0, 0] = ("charge", 24.0, clock, {"Voltage_measured": np.array([3.8, 3.9, 4.0]),
                                         "Current_measured": np.array([1.5, 1.5, 1.4]), "Time": np.array([0, 10, 20.0])})
    cyc[0, 1] = ("discharge", 24.0, clock, {"Voltage_measured": np.array([4.0, 3.7]), "Time": np.array([0, 10.0]),
                                            "Capacity": 1.85})
    cyc[0, 2] = ("impedance", 24.0, clock, {"Battery_impedance": np.array([0.2 + 0.01j, 0.21 - 0.02j]),
                                            "Re": 0.05, "Rct": 0.08})
    savemat(path, {"B9999": {"cycle": cyc}})


def test_mat_nasa_battery_layout(tmp_path):
    path = tmp_path / "B9999.mat"
    _nasa_like(path)
    wb, t = tables(path)
    assert set(t) == {"cycles", "measurements", "impedance"}
    assert list(t["cycles"]["type"]) == ["charge", "discharge", "impedance"]
    assert len(t["measurements"]) == 5 and {"battery", "cycle_index", "sample_index"} <= set(t["measurements"])
    assert {"battery_impedance_re", "battery_impedance_im"} <= set(t["impedance"].columns)
    _, generic = tables(path, mat_layout="generic")
    assert "cycles" not in generic


def test_hdf5_and_mat_v73(tmp_path):
    h5py = pytest.importorskip("h5py")
    path = tmp_path / "h.h5"
    with h5py.File(path, "w") as h:
        g = h.create_group("sensors")
        g.create_dataset("pressure", data=np.arange(8.0).reshape(4, 2))
        g.create_dataset("cube", data=np.zeros((2, 2, 2)))
    wb, t = tables(path)
    assert t["sensors.pressure"].shape == (4, 2)
    assert any("cube" in u["path"] for u in wb.unsupported)
    v73 = tmp_path / "v73.mat"
    with h5py.File(v73, "w", userblock_size=512) as h:
        h.create_dataset("signal", data=np.arange(10.0).reshape(2, 5))  # MATLAB stores column-major
    with open(v73, "r+b") as f:
        f.write(b"MATLAB 7.3 MAT-file, Platform: GLNXA64, Created on: Mon Jan  1 00:00:00 2024 HDF5 schema 1.00 .")
    wb, t = tables(v73)
    assert t["signal"].shape == (5, 2)


def test_equal_length_arrays_are_aligned_only_on_request(tmp_path):
    path = tmp_path / "s.npz"
    np.savez(path, voltage=np.array([4.1, 4.0, np.nan, 4.0, 4.1, 12.5]), temp=np.array([24.0, 24.5, 25.0, 24.6, 24.4, 130.0]),
             cycle_index=np.array([1, 2, 3, 4, 5, 5]), names=np.array(["voltage", "temp", "cycle_index"]))
    wb, t = tables(path)
    assert {"voltage", "temp", "cycle_index"} == set(t) and list(t["voltage"].columns) == ["voltage"]
    # names only lists the other arrays: metadata (and evidence they belong together), not observations
    assert any(u["path"] == "names" and u["reason"].startswith("metadata") for u in wb.unsupported)
    assert any("align_arrays" in w and "'names' lists exactly these arrays" in w for w in wb.warnings)
    wb, t = tables(path, align_arrays=True)
    assert list(t["s"].columns) == ["voltage", "temp", "cycle_index"] and set(t) == {"s"}
    res = prepare(path, tmp_path / "out", source={"align_arrays": True})
    found = {(i.column, i.rule_id) for i in res.run.issues}
    assert ("cycle_index", "unique") in found  # the repeated cycle 5
    assert any(c.startswith("aligned on request") for c in res.run.wb.tables_meta["s"]["notes"])


def test_struct_array_series_long_format_and_length_mismatch(tmp_path):
    from scipy.io import savemat

    dt = [("type", "O"), ("time", "O"), ("Voltage", "O"), ("Current", "O")]
    cyc = np.empty((1, 2), dtype=dt)
    cyc[0, 0] = ("discharge", np.array([0.0, 10, 20]), np.array([4.1, 4.0, 3.9]), np.array([-1.0, -1.0, -0.9]))
    cyc[0, 1] = ("charge", np.array([0.0, 12]), np.array([3.5, 3.8]), np.array([1.0]))
    path = tmp_path / "demo.mat"
    savemat(path, {"Demo": {"cycle": cyc}})
    wb, t = tables(path)
    assert list(t["Demo.cycle"]["cycle_index"]) == [1, 2]
    series = t["Demo.cycle.series"]
    assert len(series) == 5 and {"cycle_index", "sample_index", "time", "Voltage", "Current"} <= set(series.columns)
    assert any("Demo.cycle(2): series lengths differ" in w for w in wb.warnings)


def test_hdf5_attributes_in_structure(tmp_path):
    h5py = pytest.importorskip("h5py")
    path = tmp_path / "a.h5"
    with h5py.File(path, "w") as h:
        h.attrs["dataset_name"] = "QA example"
        d = h.create_dataset("temperature_c", data=np.array([22.1, 22.4, 140.0]))
        d.attrs["units"] = "C"
    insp = inspect_file(path)
    attrs = {n["path"]: n.get("attrs") for n in insp.tree}
    assert attrs["/"] == {"dataset_name": "QA example"} and attrs["temperature_c"] == {"units": "C"}


def test_one_off_category_next_to_established_ones_is_flagged(tmp_path):
    src = tmp_path / "c.csv"
    cities = ["Pune", "Mumbai", "Delhi", "Hyderabad"] * 8 + ["Hydrabad"]
    src.write_text("id,city\n" + "".join(f"{i},{c}\n" for i, c in enumerate(cities)))
    res = prepare(src, tmp_path / "out")
    flagged = [i for i in res.run.issues if i.column == "city" and i.rule_id == "allowed"]
    assert [i.value for i in flagged] == ["Hydrabad"]
    fix = next(s for s in res.run.suggestions if s.column == "city")
    assert fix.method == "fuzzy_category" and fix.new == "Hyderabad" and fix.status == "pending"


def _messy_frame() -> pd.DataFrame:
    cities = (["Pune", "Mumbai", "Delhi", "Chennai"] * 6)[:22] + ["unknown", "bangalore"]
    return pd.DataFrame({
        "record_id": [f"C{i:03d}" for i in range(24)],
        "city": cities,
        "spend": [str(100 + i) for i in range(22)] + ["999.00", " 1,299 "],
        "visits": [str(i % 5) for i in range(23)] + ["N/A"],
    }, dtype=object)


def test_csv_and_xlsx_give_the_same_ml_dataset(tmp_path):
    df = _messy_frame()
    df.to_csv(tmp_path / "m.csv", index=False)
    df.to_excel(tmp_path / "m.xlsx", index=False, sheet_name="m")
    out = {}
    for ext in ("csv", "xlsx"):
        ex = prepare(tmp_path / f"m.{ext}", tmp_path / f"out_{ext}").exported
        out[ext] = (ex.out_dir / next(iter(ex.datasets.values()))["csv"]).read_text(encoding="utf-8")
    assert out["csv"] == out["xlsx"]
    rows = pd.read_csv(tmp_path / "out_csv" / "m_cleaned.csv", dtype=str, keep_default_na=False)
    assert rows.loc[22, "spend"] == "999" and rows.loc[23, "spend"] == "1299"
    assert rows.loc[22, "city"] == ""  # one-off 'unknown' among real cities: a placeholder, cleared


def test_unknown_kept_when_it_is_an_established_category(tmp_path):
    src = tmp_path / "k.csv"
    status = ["Active", "Closed", "Unknown"] * 10
    src.write_text("id,status\n" + "".join(f"{i},{s}\n" for i, s in enumerate(status)))
    res = prepare(src, tmp_path / "out")
    cleaned = pd.read_csv(res.exported.out_dir / "k_cleaned.csv", dtype=str, keep_default_na=False)
    assert (cleaned["status"] == "Unknown").sum() == 10
    assert not [i for i in res.run.issues if i.column == "status"]


def test_alias_hint_names_closest_value_without_changing_it(tmp_path):
    from ilovecsv.rules.engine import _closest_hint

    cities = ["Mumbai", "Delhi", "Bengaluru", "Hyderabad"]
    assert "'Bengaluru'" in _closest_hint("bangalore", cities)
    assert _closest_hint("unknown", cities) == "" and _closest_hint("Bombay", cities) == ""  # no guess
    src = tmp_path / "c.csv"
    src.write_text("id,city\n" + "".join(f"{i},{c}\n" for i, c in enumerate(cities * 8 + ["bangalore"])))
    res = prepare(src, tmp_path / "out")
    issue = next(i for i in res.run.issues if i.rule_id == "allowed")
    assert issue.value == "bangalore" and "Closest allowed value: 'Bengaluru'" in issue.message
    cleaned = pd.read_csv(res.exported.out_dir / "c_cleaned.csv", dtype=str)
    assert cleaned["city"].iloc[-1] == "bangalore"  # flagged, never changed


def test_series_context_copies_text_fields_and_named_numbers(tmp_path):
    from scipy.io import savemat

    dt = [("type", "O"), ("ambient_temperature", "O"), ("Capacity", "O"), ("V", "O")]
    cyc = np.empty((1, 2), dtype=dt)
    cyc[0, 0] = ("discharge", 24.0, 1.8, np.array([4.1, 4.0]))
    cyc[0, 1] = ("charge", 24.0, np.nan, np.array([3.5, 3.8, 4.0]))
    path = tmp_path / "b.mat"
    savemat(path, {"B": {"cycle": cyc}})
    _, t = tables(path)
    series = t["B.cycle.series"]
    assert "type" in series.columns and "Capacity" not in series.columns and "ambient_temperature" not in series
    _, t = tables(path, series_context=["ambient_temperature"])
    assert list(t["B.cycle.series"]["ambient_temperature"]) == [24.0] * 5


def test_report_findings_group_structure_missing_and_suspicious(tmp_path):
    path = tmp_path / "s.npz"
    np.savez(path, voltage=np.array([4.1, 4.0, np.nan, 4.0, 4.1, 4.05, 12.5]),
             current=np.array([1.0, 1.0, 1.0, np.nan, 1.0, 0.9, 1.0]))
    res = prepare(path, tmp_path / "out", source={"align_arrays": True})
    findings = json.loads(res.report.read_text())["findings"]
    cats = {f["category"] for f in findings}
    assert {"missing values", "suspicious readings"} <= cats
    assert any(f["category"] == "suspicious readings" and "12.5" in f["detail"] for f in findings)
    import openpyxl

    assert "Findings" in openpyxl.load_workbook(res.exported.reports["xlsx"]).sheetnames


# ----------------------------------------------------------------------------- large files
def test_row_limit_and_sample(tmp_path):
    path = tmp_path / "big.csv"
    path.write_text("id,v\n" + "".join(f"{i},{i * 2}\n" for i in range(50)))
    with pytest.raises(LargeDatasetError):
        read_workbook(path, options=SourceOptions(row_limit=10))
    wb, t = tables(path, max_rows=10, row_limit=10)
    assert len(t["big"]) == 10 and wb.partial
    insp = inspect_file(path, options=SourceOptions(row_limit=10))
    assert not insp.full_processing_ok and insp.tables[0].rows == 50


def test_parquet_sample_marked_partial(tmp_path):
    path = tmp_path / "p.parquet"
    pd.DataFrame({"x": range(100)}).to_parquet(path)
    wb, t = tables(path, max_rows=7)
    assert len(t["p"]) == 7 and wb.partial


# ----------------------------------------------------------------------------- prepare on non-tabular sources
def test_prepare_json_handoff_files(tmp_path):
    src = tmp_path / "readings.json"
    src.write_text(json.dumps({"records": [
        {"sensor": f"S{i % 3}", "reading": " 1.5e-05 " if i == 4 else str(10 + i), "status": "ok" if i % 2 else "OK"}
        for i in range(30)]}))
    res = prepare(src, tmp_path / "out", source={"json_records": "records"})
    ex = res.exported
    assert ex.verification["passed"], ex.verification
    manifest = json.loads(ex.manifest_path.read_text())
    assert manifest["source"]["format"] == "json" and manifest["source"]["options"]["json_records"] == "records"
    assert manifest["tables"][0]["files"]["csv"]["sha256"]
    schema = json.loads(ex.schema_path.read_text())
    assert schema["schema_version"] == "1.0"
    assert "prepare_dataset.py" == ex.script_path.name and "json_records" in ex.config_path.read_text()
    cleaned = pd.read_csv(ex.out_dir / next(iter(ex.datasets.values()))["csv"])
    assert pd.to_numeric(cleaned["reading"]).notna().all()  # exponent and padded numbers parsed
    recipe = yaml.safe_load(ex.config_path.read_text())
    assert recipe["source"]["json_records"] == "records"
    again = prepare(src, tmp_path / "out2", config=ex.config_path)  # the recipe's source options are reused
    assert again.exported.datasets.keys() == ex.datasets.keys()


@pytest.mark.parametrize("name,sep", [("t.tsv", "\t"), ("semi.csv", ";")])
def test_non_comma_text_keeps_format_and_dataset_is_comma_csv(tmp_path, name, sep):
    src = tmp_path / name
    src.write_text(sep.join(["id", "city", "v"]) + "\n" + "".join(
        sep.join([str(i), " Pune " if i == 3 else "Pune", str(i)]) + "\n" for i in range(12)))
    ex = prepare(src, tmp_path / "out").exported
    assert ex.corrected_path.read_text().splitlines()[0] == sep.join(["id", "city", "v"])
    dataset = ex.out_dir / next(iter(ex.datasets.values()))["csv"]
    assert dataset != ex.corrected_path and dataset.read_text().splitlines()[0] == "id,city,v"
    assert pd.read_csv(dataset)["city"].eq("Pune").all()


def test_prepare_partial_is_labelled(tmp_path):
    src = tmp_path / "rows.csv"
    src.write_text("id,v\n" + "".join(f"{i},{i}\n" for i in range(40)))
    res = prepare(src, tmp_path / "out", source={"max_rows": 15})
    manifest = json.loads(res.exported.manifest_path.read_text())
    assert manifest["source"]["partial"]
    assert manifest["tables"][0]["rows_read"] == 15


# ----------------------------------------------------------------------------- shared engine details
def test_numeric_null_codes():
    from ilovecsv.values import is_null, normalise_nulls

    nulls = normalise_nulls(["-999", "n/a"])
    assert is_null(-999, nulls) and is_null(-999.0, nulls) and is_null("-999", nulls) and is_null(" N/A ", nulls)
    assert not is_null(-998, nulls) and not is_null(0, nulls)


def _suggested_max(values):
    from ilovecsv.profiling import profile_workbook
    from ilovecsv.reader import WorkbookData, sheet_from_frame
    from ilovecsv.rules.suggest import suggest_rules

    wb = WorkbookData(filename="x.csv", path=None, file_type="frames", sha256="",
                      sheets={"x": sheet_from_frame("x", pd.DataFrame({"v": values}))})
    return yaml.safe_load(suggest_rules(profile_workbook(wb)))["columns"]["v"].get("max")


def test_suggested_bound_dropped_when_data_breaks_it():
    rng = np.random.default_rng(0)
    bulk = list(rng.normal(50, 5, 1000).round(2))
    assert _suggested_max(bulk + [500.0, 600.0]) is not None  # two stray values: a real validity bound
    assert _suggested_max(bulk + list(rng.uniform(200, 900, 80).round(2))) is None  # 7% beyond: a heavy tail


def test_example_recipes_are_valid():
    from pathlib import Path

    from ilovecsv.rules.schema import load_rules

    for path in (Path(__file__).parents[1] / "examples" / "recipes").glob("*.yaml"):
        rules = load_rules(path.read_text(encoding="utf-8"))
        assert rules.sheets, path


def test_nasa_recipe_on_synthetic_battery(tmp_path):
    from pathlib import Path

    path = tmp_path / "B9999.mat"
    _nasa_like(path)
    recipe = Path(__file__).parents[1] / "examples" / "recipes" / "nasa_battery.yaml"
    res = prepare(path, tmp_path / "out", config=recipe)
    assert set(res.exported.datasets) == {"cycles", "impedance"}
    assert res.exported.verification["passed"]
    manifest = json.loads(res.exported.manifest_path.read_text())
    assert "battery" in json.dumps(manifest["split_hints"])  # the declared group column


# ----------------------------------------------------------------------------- CLI
def test_cli_inspect(tmp_path, capsys):
    path = tmp_path / "t.parquet"
    FRAME.to_parquet(path)
    assert cli(["inspect", str(path), "--json"]) == 0
    out = json.loads(capsys.readouterr().out)
    assert out["format"] == "parquet" and out["tables"][0]["rows"] == 4
    assert cli(["inspect", str(path)]) == 0
    assert "t" in capsys.readouterr().out
