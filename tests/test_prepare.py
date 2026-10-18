"""One-click dataset preparation for ML: Safe Auto-Clean, review queue, target checks, exports."""

import csv
import json

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from ilovecsv.api.deps import create_user
from ilovecsv.api.main import create_app
from ilovecsv.api.settings import Settings
from ilovecsv.cli import main
from ilovecsv.corrections import SAFE_METHODS
from ilovecsv.ml import load_prepared
from ilovecsv.pipeline import decide, export, prepare
from ilovecsv.reader import sha256_file
from ilovecsv.synth import build_churn_workbook


@pytest.fixture(scope="module")
def churn(tmp_path_factory):
    d = tmp_path_factory.mktemp("churn")
    path = d / "churn.xlsx"
    truth = build_churn_workbook(path)
    return path, truth


@pytest.fixture(scope="module")
def prepared(churn, tmp_path_factory):
    path, truth = churn
    res = prepare(path, tmp_path_factory.mktemp("out"), target="churn", formats=("csv", "parquet"))
    report = json.loads(res.report.read_text(encoding="utf-8"))
    return res, report, truth


def _row_of(res, cid):
    df = res.run.wb.sheets["Customers"].df
    return int(df.index[df["customer_id"] == cid][0])


def cleaned_df(res):
    return pd.read_csv(res.dataset_path("csv"), keep_default_na=False, na_values=[""])


# ----------------------------------------------------------------------------- safe vs review


def test_safe_auto_clean_applies_only_safe_fixes(prepared):
    res, report, _ = prepared
    applied = [c for c in report["changes"]]
    assert applied and all(c["safety"] == "safe" and c["method"] in SAFE_METHODS for c in applied)
    for s in res.run.suggestions:
        if s.safety == "safe":
            assert s.status == "approved", s
        else:
            assert s.status == "pending", s
    assert report["config"]["mode"] == "safe_auto_clean"


def test_duplicates_wait_for_review_by_default(prepared):
    res, report, truth = prepared
    queued = [r for r in report["review_required"] if r.get("action") == "delete_row"]
    assert len(queued) == 2  # the exact copy and the near-duplicate
    assert report["removed_rows"] == []
    assert report["dataset"]["rows_before"] == report["dataset"]["rows_after"] == {"Customers": 403}


def test_exact_duplicate_policy_removes_only_exact_copies(churn, tmp_path):
    path, truth = churn
    res = prepare(path, tmp_path, target="churn", remove_exact_duplicates=True)
    report = json.loads(res.report.read_text(encoding="utf-8"))
    assert [r["exact_copy"] for r in report["removed_rows"]] == [True]
    assert report["removed_rows"][0]["values"]["customer_id"] == truth["exact_duplicate"]
    near = [r for r in report["review_required"] if r.get("action") == "delete_row"]
    assert len(near) == 1  # the near-duplicate (different id) still needs a person
    assert report["dataset"]["rows_after"] == {"Customers": 402}


# ----------------------------------------------------------------------------- cleaning results


def test_null_placeholders_become_empty_cells(prepared):
    res, report, truth = prepared
    df = cleaned_df(res)
    for cid, col in truth["null_tokens"]:
        assert pd.isna(df.loc[df.customer_id == cid, col].iloc[0]), (cid, col)
    raw = list(csv.reader(open(res.dataset_path("csv"), encoding="utf-8")))
    assert not any(cell in ("N/A", "n/a", "-", "null") for row in raw for cell in row)


def test_labels_normalised_but_not_rewritten(prepared):
    res, report, truth = prepared
    df = cleaned_df(res)
    assert set(df["churn"].dropna()) == {"Yes", "No"}  # 'yes', 'NO ', 'YES' fixed by case/spacing only
    labelled = df.dropna(subset=["churn"])
    assert len(labelled) == 401  # nothing dropped: missing labels are flagged, not removed


def test_numbers_as_text_converted(prepared):
    res, _, truth = prepared
    df = cleaned_df(res)
    for cid in truth["numbers_as_text"]:
        assert df.loc[df.customer_id == cid, "monthly_charges"].iloc[0] > 1000


def test_target_checks(prepared):
    res, report, truth = prepared
    ml = report["ml"]["sheets"][0]
    assert ml["target"] == "churn"
    assert sorted(ml["missing_target_rows"]) == sorted(_row_of(res, c) for c in truth["missing_label"])
    assert [sorted(g) for g in ml["label_conflicts"]] == [sorted(_row_of(res, c) for c in truth["conflict"])]
    assert {c["label"] for c in ml["target_summary"]["classes"]} == {"Yes", "No"}
    assert [s["column"] for s in ml["leakage_suspects"]] == ["cancellation_date"]
    assert not any(a.column == "churn" for a in res.run.anomalies)  # rare labels are not anomalies


def test_feature_schema(prepared):
    _, report, _ = prepared
    sch = report["columns"][0]
    assert sch["features"]["numeric"] == ["monthly_charges", "tenure_months", "support_calls"]
    assert sch["features"]["categorical"] == ["plan", "contract"]
    assert sch["features"]["datetime"] == ["signup_date"]  # cancellation_date leaks, so it is not "ready"
    assert [r["column"] for r in sch["needs_review"]] == ["cancellation_date"]
    assert sch["excluded"] == {"identifier": ["customer_id"], "pii": ["email"], "empty": [], "target": "churn"}


def test_load_prepared_and_parquet_dtypes(prepared):
    res, _, _ = prepared
    df, schema = load_prepared(res.report)
    assert str(df["tenure_months"].dtype) == "Int64"
    assert str(df["monthly_charges"].dtype) == "Float64"
    assert str(df["plan"].dtype) == "category"
    assert pd.api.types.is_datetime64_any_dtype(df["signup_date"])
    pq = pd.read_parquet(res.dataset_path("parquet"))
    assert str(pq["support_calls"].dtype) == "Int64" and str(pq["churn"].dtype) == "category"


def test_unparseable_column_kept_as_text(churn, tmp_path):
    from openpyxl import load_workbook

    path, _ = churn
    wb = load_workbook(path)
    wb["Customers"]["F10"] = "about a year"  # tenure_months
    bad = tmp_path / "bad.xlsx"
    wb.save(bad)
    res = prepare(bad, tmp_path / "out", target="churn", formats=("csv", "parquet"))
    kept = res.exported.kept_as_text
    assert [(k["column"], k["unparseable_values"]) for k in kept] == [("tenure_months", 1)]
    assert str(pd.read_parquet(res.dataset_path("parquet"))["tenure_months"].dtype).startswith("str")
    assert any(u["rule_id"] == "type" for u in res.unresolved)  # and it is reported as unresolved


# ----------------------------------------------------------------------------- reproducibility & report


def test_outputs_original_and_config(prepared, churn):
    res, report, _ = prepared
    out = res.exported.out_dir
    assert sha256_file(out / report["input"]["preserved_copy"]) == sha256_file(churn[0])
    assert report["dataset"]["config_file"] == "churn_prepare_config.yaml"
    names = {p.name for p in out.iterdir()}
    assert {"churn_cleaned.csv", "churn_cleaned.parquet", "churn_cleaned.xlsx", "churn_cleaning_report.json",
            "churn_cleaning_report.xlsx", "churn_cleaning_report.html", "churn_prepare_config.yaml"} <= names


def test_saved_config_cleans_new_batch_identically(prepared, tmp_path):
    res, report, _ = prepared
    batch = tmp_path / "batch2.csv"
    build_churn_workbook(batch, n=150, seed=99, as_csv=True)
    res2 = prepare(batch, tmp_path / "out2", config=res.exported.config_path)
    assert res2.run.rules_sha256 == res.run.rules_sha256  # same rules, not re-suggested from new data
    assert res2.run.rules.sheets[0].target == "churn"


def test_report_contract(prepared):
    _, report, _ = prepared
    assert report["schema_version"] == "1.0"
    for key in ("tool", "run", "input", "config", "dataset", "columns", "ml", "summary", "changes",
                "removed_rows", "review_required", "unresolved", "issues"):
        assert key in report, key
    change = report["changes"][0]
    assert {"row", "row_cleaned", "column", "old", "new", "method", "safety", "approved_by"} <= change.keys()
    assert all("row_original" in u and "row_cleaned" in u for u in report["unresolved"])


def test_manual_flow_still_has_no_auto_changes(churn, tmp_path):
    from ilovecsv.pipeline import run_on_workbook
    from ilovecsv.reader import read_workbook
    from ilovecsv.rules.schema import load_rules

    path, _ = churn
    rules = "sheet: Customers\ntarget: churn\n"
    run = run_on_workbook(read_workbook(path), load_rules(rules), rules)
    ex = export(run, tmp_path)
    assert ex.apply.applied == []  # nothing approved -> unchanged copy
    decide(run, mode="safe")
    assert run.auto_cleaned and all(s.status == "approved" for s in run.suggestions if s.safety == "safe")


# ----------------------------------------------------------------------------- entry points


def test_cli_prepare(churn, tmp_path, capsys):
    path, _ = churn
    code = main(["prepare", str(path), "--target", "churn", "--out", str(tmp_path)])
    out = capsys.readouterr().out
    assert code == 1  # missing labels are unresolved errors
    assert "Safe Auto-Clean applied" in out and "cancellation_date: possible target leakage" in out
    assert (tmp_path / "churn_cleaned.csv").exists()
    assert main(["prepare", str(path), "--target", "nope", "--out", str(tmp_path / "x")]) == 2


def test_api_prepare_and_downloads(churn, tmp_path):
    path, _ = churn
    app = create_app(Settings(data_dir=tmp_path / "data"))
    with app.state.sessionmaker() as db:
        key = {"X-API-Key": create_user(db, "ml")}
    c = TestClient(app, raise_server_exceptions=False)
    with open(path, "rb") as f:
        wid = c.post("/v1/workbooks", headers=key, files={"file": ("churn.xlsx", f)}).json()["id"]
    bad = c.post(f"/v1/workbooks/{wid}/prepare", headers=key, json={"target": "nope"})
    assert bad.status_code == 422 and bad.json()["type"] == "/problems/invalid-target"
    r = c.post(f"/v1/workbooks/{wid}/prepare", headers=key, json={"target": "churn", "formats": ["csv", "parquet"]})
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["mode"] == "safe_auto_clean" and p["status"] == "exported"
    assert p["safe_fixes_applied"] > 0 and p["review_required"] >= 2
    csv_r = c.get(f"/v1/runs/{p['id']}/dataset", headers=key, params={"format": "csv"})
    assert csv_r.status_code == 200 and csv_r.content.startswith(b"customer_id,")
    assert c.get(f"/v1/runs/{p['id']}/dataset", headers=key, params={"format": "parquet"}).content[:4] == b"PAR1"
    assert b"target: churn" in c.get(f"/v1/runs/{p['id']}/config", headers=key).content
    report = c.get(f"/v1/runs/{p['id']}/report", headers=key, params={"format": "json"}).json()
    assert report["schema_version"] == "1.0" and report["ml"]["sheets"][0]["leakage_suspects"]
    sugs = c.get(f"/v1/runs/{p['id']}/suggestions", headers=key).json()["items"]
    assert {s["safety"] for s in sugs if s["status"] == "approved"} == {"safe"}
    # approve_safe on the manual flow does the same thing
    run2 = c.post(f"/v1/workbooks/{wid}/runs", headers=key, json={}).json()
    d = c.post(f"/v1/runs/{run2['id']}/decisions", headers=key, json={"approve_safe": True}).json()
    assert d["changed"] > 0
