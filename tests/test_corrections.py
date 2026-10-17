"""Safe correction workflow: only approved changes, nothing else touched."""

from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from ilovecsv.corrections import closest_category
from ilovecsv.models import APPROVED, PENDING
from ilovecsv.pipeline import decide, export, run_checks


def by_method(run):
    return {s.method: s for s in run.suggestions}


def test_expected_suggestions(orders_run):
    got = {(s.cell_ref, s.method, str(s.new)) for s in orders_run.suggestions}
    methods = {m for _, m, _ in got}
    assert {"trim_whitespace", "canonical_category", "fuzzy_category", "text_to_number", "text_to_date",
            "serial_to_date", "remove_duplicate_row", "rule_fix"} <= methods
    news = {(m, n) for _, m, n in got}
    assert ("fuzzy_category", "Shipped") in news
    assert ("fuzzy_category", "North") in news
    assert ("canonical_category", "Delivered") in news
    assert ("text_to_date", "2024-03-15") in news
    assert ("serial_to_date", "2023-03-15") in news
    assert ("text_to_number", "1200.5") in news
    # 'Lost' has no close allowed value, so no guess is made.
    assert not any(s.old == "Lost" for s in orders_run.suggestions)
    # Missing required values are never imputed.
    required_ids = {i.id for i in orders_run.validation.issues if i.rule_id == "required"}
    assert not any(s.issue_id in required_ids for s in orders_run.suggestions)
    assert all(s.status == PENDING for s in orders_run.suggestions)


def test_closest_category():
    assert closest_category("Shiped", ["Pending", "Shipped"])[0] == "Shipped"
    assert closest_category("Lost", ["Pending", "Shipped"]) is None
    assert closest_category("Nort", ["North", "South"])[0] == "North"


@pytest.fixture()
def fresh_run(orders, rules_text):
    return run_checks(orders[0], rules_text)


def test_only_approved_are_applied(fresh_run, tmp_path):
    run = fresh_run
    m = by_method(run)
    approve = [m["fuzzy_category"].id, m["text_to_date"].id]
    reject = [m["trim_whitespace"].id]
    decide(run, approve=approve, reject=reject, approver="alice")
    ex = export(run, tmp_path)
    assert {c["suggestion_id"] for c in ex.apply.applied} == set(approve)

    original = load_workbook(run.wb.path)
    corrected = load_workbook(ex.corrected_path)
    changed_cells = {c["cell_ref"] for c in ex.apply.applied}
    for ws in original.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                ref = f"{ws.title}!{cell.coordinate}"
                new_val = corrected[ws.title][cell.coordinate].value
                if ref in changed_cells:
                    assert new_val != cell.value
                else:
                    assert new_val == cell.value, f"{ref} changed without approval"

    date_cell = corrected[m["text_to_date"].sheet][m["text_to_date"].cell_ref.split("!")[1]]
    assert isinstance(date_cell.value, datetime) and date_cell.number_format == "yyyy-mm-dd"
    assert "was: '15/03/2024'" in date_cell.comment.text and "alice" in date_cell.comment.text
    assert date_cell.fill.fgColor.rgb.endswith("FFF2CC")
    assert corrected["Summary"]["B2"].value == "=SUM(Orders!J:J)"  # formulas preserved
    assert corrected["Orders"].freeze_panes == "A2"  # sheet settings preserved


def test_revalidation_shows_effect(fresh_run, tmp_path):
    run = fresh_run
    decide(run, min_confidence=0.0, approver="bot")
    ex = export(run, tmp_path)
    before, after = run.summary(), ex.after
    assert after["warnings"] < before["warnings"]
    for fixed in ("type.date_as_text", "type.date_serial", "type.stored_as_text", "format.whitespace",
                  "allowed.noncanonical", "total_matches"):
        assert fixed not in after["by_rule"], fixed
    # Fixing 'Shiped' -> 'Shipped' exposes the missing ship date: re-validation must report it.
    assert after["by_rule"].get("shipped_needs_date", 0) == before["by_rule"]["shipped_needs_date"] + 1


def test_duplicate_row_kept_when_workbook_has_formulas(fresh_run, tmp_path):
    run = fresh_run
    dup = by_method(run)["remove_duplicate_row"]
    decide(run, approve=[dup.id])
    ex = export(run, tmp_path)
    assert [s["suggestion_id"] for s in ex.apply.skipped] == [dup.id]
    assert "formulas" in ex.apply.skipped[0]["reason"]


def test_duplicate_row_deleted_without_formulas(tmp_path, rules_text):
    from ilovecsv.synth import build_orders_workbook

    path = tmp_path / "nf.xlsx"
    build_orders_workbook(path, with_formula_sheet=False)
    run = run_checks(path, rules_text)
    dup = by_method(run)["remove_duplicate_row"]
    decide(run, approve=[dup.id])
    ex = export(run, tmp_path / "out")
    assert [a["suggestion_id"] for a in ex.apply.applied] == [dup.id]
    assert "row.duplicate" not in ex.after["by_rule"]
    assert load_workbook(ex.corrected_path)["Orders"].max_row == load_workbook(path)["Orders"].max_row - 1


def test_formula_cells_never_overwritten(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["a", "b"])
    ws.append([" x ", "=A2"])
    p = tmp_path / "f.xlsx"
    wb.save(p)
    run = run_checks(p, "sheet: S\ncolumns:\n  b: {required: true}\n")
    assert not any(i.column == "b" for i in run.validation.issues)  # uncomputed formula is not "missing"
    decide(run, min_confidence=0)
    ex = export(run, tmp_path / "out")
    out = load_workbook(ex.corrected_path)["S"]
    assert out["A2"].value == "x" and out["B2"].value == "=A2"


def test_stale_cell_is_skipped(fresh_run, tmp_path):
    run = fresh_run
    s = by_method(run)["trim_whitespace"]
    s.status = APPROVED
    s.old = "something else"  # simulate the file changing after validation
    ex = export(run, tmp_path)
    assert any(x["suggestion_id"] == s.id and "changed" in x["reason"] for x in ex.apply.skipped)


def test_csv_corrections(tmp_path):
    p = tmp_path / "c.csv"
    p.write_text("id,status,qty\n1,Pending ,5\n2,shipped,\"1,200\"\n", encoding="utf-8")
    rules = "columns:\n  status: {allowed: [Pending, Shipped]}\n  qty: {type: int}\n"
    run = run_checks(p, rules)
    decide(run, min_confidence=0.9)
    ex = export(run, tmp_path / "out")
    assert ex.corrected_path.suffix == ".csv"
    text = ex.corrected_path.read_text(encoding="utf-8").splitlines()
    assert text == ["id,status,qty", "1,Pending,5", "2,Shipped,1200"]
    assert ex.after["errors"] == 0 and ex.after["warnings"] == 0
